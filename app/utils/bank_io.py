"""Whole-question-bank import / export.

One spreadsheet describes many tests at once (a "Test Title" column says which
test each question belongs to). Importing is deliberately careful:

* It works out a *plan* first (nothing is written) so the teacher can preview
  exactly what will be created, changed, or left alone.
* Applying the plan UPDATES existing questions in place, matched by test title
  and question number. It never deletes and re-creates a question that
  students have already answered.
* If the options of an already-answered question are re-ordered, students'
  recorded answers are moved to follow their option text, so a student who
  chose "Unlimited" still has "Unlimited" after it moves from B to C.
* Tests that students have already taken are re-graded afterwards.
"""
import io
import logging
import re
import secrets
import threading
import time
from collections import defaultdict

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from sqlalchemy import func

from app.extensions import db
from app.models import Question, Response, Submission, Test
from app.utils.grading import recalculate_test

logger = logging.getLogger(__name__)

BANK_HEADERS = [
    "Test Title", "Test Number", "Chapter", "Subject", "Q No", "Question",
    "Option A", "Option B", "Option C", "Option D", "Correct Answer",
    "Marks", "Negative Marks", "Explanation",
]

_ALIASES = {
    "title": ("test title", "test", "test name", "title"),
    "test_number": ("test number", "test no", "test no."),
    "chapter": ("chapter", "unit", "chapter / unit", "chapter/unit"),
    "subject": ("subject",),
    "q_number": ("q no", "q. no.", "q no.", "q.no.", "question no", "question number", "qno"),
    "text": ("question", "question text"),
    "option_a": ("option a", "opt a"),
    "option_b": ("option b", "opt b"),
    "option_c": ("option c", "opt c"),
    "option_d": ("option d", "opt d"),
    "correct_answer": ("correct answer", "answer", "key", "answer key"),
    "marks": ("marks", "mark"),
    "negative_marks": ("negative marks", "negative mark", "negative"),
    "explanation": ("explanation", "solution"),
}

MAX_ERRORS = 200
_OPTION_FIELDS = ("option_a", "option_b", "option_c", "option_d")
_GRADING_FIELDS = ("correct_answer", "marks", "negative_marks") + _OPTION_FIELDS


def _clean(value):
    return "" if value is None else str(value).strip()


def _norm(text):
    """Comparison form of an option's text (whitespace, case, trailing dot)."""
    return re.sub(r"\s+", " ", _clean(text)).casefold().rstrip(".")


# ------------------------------------------------------------------ parsing

def parse_bank(file_stream):
    """Returns (rows, errors, warnings). Nothing touches the database."""
    try:
        wb = load_workbook(file_stream, read_only=True, data_only=True)
    except Exception as exc:
        return [], [f"Could not read the file: {exc}"], []

    ws = wb.worksheets[0]
    iterator = ws.iter_rows(values_only=True)
    try:
        header = [(_clean(h).casefold()) for h in next(iterator)]
    except StopIteration:
        return [], ["The file is empty."], []

    col = {}
    for field, names in _ALIASES.items():
        for name in names:
            if name in header:
                col[field] = header.index(name)
                break

    errors, warnings = [], []
    for required, label in (("title", "Test Title"), ("correct_answer", "Correct Answer")):
        if required not in col:
            errors.append(f"Missing required column: '{label}'.")
    if "option_a" not in col or "option_b" not in col:
        errors.append("Missing option columns: 'Option A' to 'Option D'.")
    if errors:
        return [], errors, []

    rows = []
    seen = set()
    counters = defaultdict(int)
    for row_no, values in enumerate(iterator, start=2):
        if all(v is None or _clean(v) == "" for v in values):
            continue

        def get(field):
            idx = col.get(field)
            if idx is None or idx >= len(values):
                return None
            return values[idx]

        problems = []
        title = _clean(get("title"))
        if not title:
            problems.append("Test Title is empty")
        elif len(title) > 255:
            problems.append("Test Title is longer than 255 characters")

        raw_q = get("q_number")
        if "q_number" in col:
            try:
                q_number = int(float(_clean(raw_q)))
            except ValueError:
                q_number = None
                problems.append(f"Q No '{_clean(raw_q)}' is not a number")
        else:
            counters[title.casefold()] += 1
            q_number = counters[title.casefold()]

        correct = _clean(get("correct_answer")).upper()
        if correct not in ("A", "B", "C", "D"):
            problems.append(f"Correct Answer must be A, B, C or D (found '{correct}')")

        options = {f: _clean(get(f)) for f in _OPTION_FIELDS}
        filled = [f for f, v in options.items() if v]
        if len(filled) < 2:
            problems.append("needs at least two options filled in")
        if correct in ("A", "B", "C", "D") and not options[f"option_{correct.lower()}"]:
            problems.append(f"the Correct Answer ({correct}) points to an empty option")
        texts = [_norm(options[f]) for f in filled]
        if len(texts) != len(set(texts)):
            warnings.append(f"Row {row_no}: two options have identical text.")

        marks = neg = None
        for field in ("marks", "negative_marks"):
            raw = _clean(get(field))
            if raw:
                try:
                    number = float(raw)
                except ValueError:
                    problems.append(f"{field.replace('_', ' ')} '{raw}' is not a number")
                    continue
                if field == "marks":
                    marks = number
                else:
                    neg = abs(number)

        chapter = _clean(get("chapter"))
        test_number = _clean(get("test_number"))
        subject = _clean(get("subject"))
        if len(chapter) > 150:
            problems.append("Chapter is longer than 150 characters")
        if len(test_number) > 50:
            problems.append("Test Number is longer than 50 characters")
        if len(subject) > 120:
            problems.append("Subject is longer than 120 characters")

        if title and q_number is not None:
            key = (title.casefold(), q_number)
            if key in seen:
                problems.append(f"'{title}' Q{q_number} appears twice in the file")
            seen.add(key)

        if problems:
            if len(errors) < MAX_ERRORS:
                errors.append(f"Row {row_no}: " + "; ".join(problems) + ".")
            continue

        explanation = _clean(get("explanation")) if "explanation" in col else ""
        rows.append({
            "row_no": row_no, "title": title, "test_number": test_number,
            "chapter": chapter, "subject": subject, "q_number": q_number,
            "text": _clean(get("text")), **options, "correct_answer": correct,
            "marks": marks, "negative_marks": neg,
            "explanation": explanation or None,
        })

    if not rows and not errors:
        errors.append("No question rows found.")
    if len(errors) >= MAX_ERRORS:
        errors.append(f"...stopped after {MAX_ERRORS} problems; fix these and upload again.")
    return rows, errors, warnings


# ----------------------------------------------------------------- planning

def _letter_map(old_q, new_row):
    """{old letter: new letter} following each option's text. Only letters whose
    text appears exactly once among the new options are mapped."""
    new_by_text = defaultdict(list)
    for letter in "ABCD":
        text = _norm(new_row[f"option_{letter.lower()}"])
        if text:
            new_by_text[text].append(letter)
    mapping, unmapped = {}, []
    for letter in "ABCD":
        text = _norm(old_q.option_text(letter))
        if not text:
            continue
        hits = new_by_text.get(text, [])
        if len(hits) == 1:
            mapping[letter] = hits[0]
        else:
            unmapped.append(letter)
    return mapping, unmapped


def build_plan(teacher_id, rows):
    """Compare the file against the database. Read-only."""
    by_title = defaultdict(list)
    for r in rows:
        by_title[r["title"].casefold()].append(r)

    existing = {}
    for t in Test.query.filter_by(teacher_id=teacher_id).order_by(Test.id):
        existing.setdefault(t.title.strip().casefold(), t)

    matched_ids = [existing[k].id for k in by_title if k in existing]
    questions = defaultdict(dict)
    duplicate_warnings = []
    if matched_ids:
        for q in Question.query.filter(Question.test_id.in_(matched_ids)).order_by(Question.id):
            if q.q_number in questions[q.test_id]:
                duplicate_warnings.append(
                    f"'{existing_by_id(existing, q.test_id).title}' already has two questions numbered {q.q_number}; only the first is updated.")
                continue
            questions[q.test_id][q.q_number] = q
    sub_counts = {}
    if matched_ids:
        sub_counts = dict(
            db.session.query(Submission.test_id, func.count(Submission.id))
            .filter(Submission.test_id.in_(matched_ids)).group_by(Submission.test_id).all()
        )

    tests = []
    for key, group in by_title.items():
        group.sort(key=lambda r: (r["q_number"], r["row_no"]))
        test = existing.get(key)
        first = group[0]
        tp = {
            "teacher_id": teacher_id, "title": first["title"], "test": test,
            "test_id": test.id if test else None, "exists": test is not None,
            "test_number": next((r["test_number"] for r in group if r["test_number"]), ""),
            "chapter": next((r["chapter"] for r in group if r["chapter"]), ""),
            "subject": next((r["subject"] for r in group if r["subject"]), ""),
            "default_marks": next((r["marks"] for r in group if r["marks"] is not None), 1.0),
            "default_neg": next((r["negative_marks"] for r in group if r["negative_marks"] is not None), 0.0),
            "submissions": sub_counts.get(test.id, 0) if test else 0,
            "questions": [], "added": 0, "updated": 0, "same": 0, "missing": [],
            "meta_changes": [], "grading_changes": False,
        }
        if test:
            for field, label in (("test_number", "Test Number"), ("chapter", "Chapter"), ("subject", "Subject")):
                new = tp[field]
                old = _clean(getattr(test, field))
                if new and new != old:
                    tp["meta_changes"].append((label, old, new))
            db_questions = questions.get(test.id, {})
            tp["missing"] = sorted(set(db_questions) - {r["q_number"] for r in group})
        else:
            db_questions = {}

        for r in group:
            old = db_questions.get(r["q_number"])
            qp = {"row": r, "q_number": r["q_number"], "question": old, "changes": [],
                  "mapping": {}, "unmapped": [], "action": "add"}
            if old is not None:
                for field in ("text",) + _OPTION_FIELDS + ("correct_answer",):
                    if _clean(getattr(old, field)) != r[field]:
                        qp["changes"].append((field, _clean(getattr(old, field)), r[field]))
                for field in ("marks", "negative_marks"):
                    if r[field] is not None and float(getattr(old, field) or 0) != r[field]:
                        qp["changes"].append((field, getattr(old, field), r[field]))
                if r["explanation"] is not None and _clean(old.explanation) != r["explanation"]:
                    qp["changes"].append(("explanation", _clean(old.explanation), r["explanation"]))
                qp["action"] = "update" if qp["changes"] else "same"
                if any(c[0] in _OPTION_FIELDS for c in qp["changes"]) and tp["submissions"]:
                    qp["mapping"], qp["unmapped"] = _letter_map(old, r)
                if any(c[0] in _GRADING_FIELDS for c in qp["changes"]):
                    tp["grading_changes"] = True
            tp[{"add": "added", "update": "updated", "same": "same"}[qp["action"]]] += 1
            tp["questions"].append(qp)
        tests.append(tp)

    return {"tests": tests, "warnings": sorted(set(duplicate_warnings))}


def existing_by_id(existing, test_id):
    for t in existing.values():
        if t.id == test_id:
            return t
    return None


def _count_moves(question_id, letters):
    if not letters:
        return 0
    return (
        db.session.query(func.count(Response.id))
        .filter(Response.question_id == question_id, Response.selected_answer.in_(list(letters)))
        .scalar() or 0
    )


def plan_summary(plan):
    """Plain numbers + short lists for the preview page."""
    s = {"tests_new": 0, "tests_existing": 0, "added": 0, "updated": 0, "same": 0,
         "tests_regraded": 0, "answers_moved": 0, "answers_unmapped": 0, "attempted_changed": 0}
    per_test, samples = [], []
    for tp in plan["tests"]:
        s["tests_new" if not tp["exists"] else "tests_existing"] += 1
        s["added"] += tp["added"]; s["updated"] += tp["updated"]; s["same"] += tp["same"]
        moved = unmapped = 0
        for qp in tp["questions"]:
            if qp["question"] is None:
                continue
            moved += _count_moves(qp["question"].id, [o for o, n in qp["mapping"].items() if o != n])
            unmapped += _count_moves(qp["question"].id, qp["unmapped"])
            if qp["action"] == "update" and len(samples) < 60:
                samples.append((tp["title"], qp["q_number"], qp["changes"]))
        if tp["submissions"] and tp["grading_changes"]:
            s["tests_regraded"] += 1
        s["answers_moved"] += moved; s["answers_unmapped"] += unmapped
        per_test.append({
            "title": tp["title"], "exists": tp["exists"], "chapter": tp["chapter"],
            "added": tp["added"], "updated": tp["updated"], "same": tp["same"],
            "missing": len(tp["missing"]), "submissions": tp["submissions"],
            "regrade": bool(tp["submissions"] and tp["grading_changes"]),
        })
    return s, per_test, samples


# ---------------------------------------------------------------- applying

def _remap_answers(question_id, mapping):
    """Move recorded answers to follow their option text. Done in two passes
    (through lowercase letters) so A<->B swaps cannot collide."""
    moves = {old: new for old, new in mapping.items() if old != new}
    if not moves:
        return 0
    moved = 0
    for old, new in moves.items():
        moved += Response.query.filter(
            Response.question_id == question_id, Response.selected_answer == old
        ).update({"selected_answer": new.lower()}, synchronize_session=False)
    Response.query.filter(
        Response.question_id == question_id,
        Response.selected_answer.in_([n.lower() for n in moves.values()]),
    ).update({"selected_answer": func.upper(Response.selected_answer)}, synchronize_session=False)
    return moved


def _apply_test(tp, delete_missing, stats, log):
    # Everything is re-fetched here (one query per test). Objects loaded while
    # planning are expired by the previous test's commit, and touching them
    # again would cost one extra database round trip per question.
    test = db.session.get(Test, tp["test_id"]) if tp["test_id"] else None
    current = {}
    if test is not None:
        for q in Question.query.filter_by(test_id=test.id).order_by(Question.id):
            current.setdefault(q.q_number, q)
    if test is None:
        test = Test(
            teacher_id=tp["teacher_id"], title=tp["title"],
            subject=tp["subject"] or "CA Foundation Economics",
            test_number=tp["test_number"], chapter=tp["chapter"],
            marks_per_question_default=tp["default_marks"],
            negative_marks_default=tp["default_neg"], status="draft",
        )
        db.session.add(test)
        db.session.flush()
        stats["tests_created"] += 1
    else:
        for label, old, new in tp["meta_changes"]:
            setattr(test, {"Test Number": "test_number", "Chapter": "chapter", "Subject": "subject"}[label], new)
            log.append((tp["title"], "", label, old, new))
        if tp["meta_changes"]:
            stats["tests_updated"] += 1

    for qp in tp["questions"]:
        r = qp["row"]
        if qp["action"] == "add":
            db.session.add(Question(
                test_id=test.id, order_index=r["q_number"], q_number=r["q_number"],
                text=r["text"], option_a=r["option_a"], option_b=r["option_b"],
                option_c=r["option_c"], option_d=r["option_d"],
                correct_answer=r["correct_answer"],
                marks=r["marks"] if r["marks"] is not None else test.marks_per_question_default,
                negative_marks=r["negative_marks"] if r["negative_marks"] is not None else test.negative_marks_default,
                explanation=r["explanation"] or None,
            ))
            stats["questions_added"] += 1
            log.append((tp["title"], r["q_number"], "(new question)", "", r["text"][:80]))
            continue
        if qp["action"] == "same":
            continue

        q = current.get(qp["q_number"])
        if q is None:
            continue
        for field, old, new in qp["changes"]:
            setattr(q, field, new)
            log.append((tp["title"], q.q_number, field, old, new))
        if any(c[0] in ("text", "correct_answer") + _OPTION_FIELDS for c in qp["changes"]):
            q.ai_explanation = None
        if qp["mapping"]:
            moved = _remap_answers(q.id, qp["mapping"])
            if moved:
                stats["answers_moved"] += moved
                log.append((tp["title"], q.q_number, "students' answers moved to follow option text", "", f"{moved} answer(s)"))
        if qp["unmapped"]:
            stuck = _count_moves(q.id, qp["unmapped"])
            if stuck:
                stats["answers_unmapped"] += stuck
                log.append((tp["title"], q.q_number, "WARNING: option text was edited, so these answers could not be moved automatically", "", f"{stuck} answer(s) left on their original letter"))
        stats["questions_updated"] += 1

    db.session.flush()
    if delete_missing and tp["missing"]:
        if tp["submissions"] == 0:
            for number in tp["missing"]:
                q = current.get(number)
                if q is None:
                    continue
                log.append((tp["title"], q.q_number, "(question removed - not in file)", (q.text or "")[:80], ""))
                db.session.delete(q)
                stats["questions_deleted"] += 1
        else:
            log.append((tp["title"], "", "NOT removed: questions missing from the file, but students have attempted this test", "", str(tp["missing"])))

    db.session.flush()
    ordered = Question.query.filter_by(test_id=test.id).order_by(Question.q_number, Question.id).all()
    for i, q in enumerate(ordered, start=1):
        if q.order_index != i:
            q.order_index = i
    return test.id


def apply_plan(plan, delete_missing, progress=None):
    """Returns (stats, change_log, errors). Commits one test at a time so one
    bad test cannot undo the others."""
    stats = defaultdict(int)
    log, errors = [], []
    total = len(plan["tests"])
    for done, tp in enumerate(plan["tests"], start=1):
        try:
            test_id = _apply_test(tp, delete_missing, stats, log)
            db.session.commit()
            if tp["submissions"] and tp["grading_changes"]:
                recalculate_test(db.session.get(Test, test_id))
                stats["tests_regraded"] += 1
        except Exception as exc:
            db.session.rollback()
            logger.exception("Bank import failed for test %s", tp["title"])
            errors.append(f"'{tp['title']}' was skipped: {exc}")
        if progress:
            progress(done, total)
    return dict(stats), log, errors


# --------------------------------------------- background job bookkeeping

_pending = {}      # token -> {"teacher_id", "rows", "at"}  (uploaded, awaiting confirmation)
_jobs = {}         # token -> job state
_lock = threading.Lock()


def _expire():
    now = time.time()
    for token in [t for t, p in _pending.items() if now - p["at"] > 3600]:
        _pending.pop(token, None)
    for token in [t for t, j in _jobs.items() if now - j["started"] > 6 * 3600]:
        _jobs.pop(token, None)


def stash_upload(teacher_id, rows):
    with _lock:
        _expire()
        token = secrets.token_hex(12)
        _pending[token] = {"teacher_id": teacher_id, "rows": rows, "at": time.time()}
        return token


def get_upload(token, teacher_id):
    with _lock:
        p = _pending.get(token)
        return p["rows"] if p and p["teacher_id"] == teacher_id else None


def start_job(app, token, teacher_id, delete_missing):
    with _lock:
        p = _pending.pop(token, None)
    if not p or p["teacher_id"] != teacher_id:
        return None
    _jobs[token] = {
        "teacher_id": teacher_id, "started": time.time(), "state": "running",
        "done": 0, "total": 0, "stats": {}, "errors": [], "log": [],
    }

    def run():
        job = _jobs[token]
        with app.app_context():
            try:
                plan = build_plan(teacher_id, p["rows"])
                job["total"] = len(plan["tests"])

                def progress(done, total):
                    job["done"], job["total"] = done, total

                job["stats"], job["log"], job["errors"] = apply_plan(plan, delete_missing, progress)
                job["state"] = "done"
            except Exception as exc:
                logger.exception("Bank import job failed")
                job["errors"].append(str(exc))
                job["state"] = "failed"
            finally:
                db.session.remove()

    threading.Thread(target=run, daemon=True).start()
    return token


def get_job(token, teacher_id):
    job = _jobs.get(token)
    return job if job and job["teacher_id"] == teacher_id else None


# ---------------------------------------------------------- export helpers

def _style(ws, headers, widths):
    for col, title in enumerate(headers, start=1):
        cell = ws.cell(row=1, column=col, value=title)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", start_color="1E3A5F", end_color="1E3A5F")
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for col, width in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(col)].width = width
    ws.freeze_panes = "A2"


def _to_bytes(wb):
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf


_BANK_WIDTHS = (28, 12, 22, 20, 8, 55, 25, 25, 25, 25, 10, 8, 10, 70)


def export_bank(teacher_id):
    """Every question of every test this teacher owns, in the import format.
    Doubles as a backup: re-importing the untouched file changes nothing."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Question Bank"
    headers = BANK_HEADERS + ["Status (ignored on import)", "Old AI explanation (ignored on import)"]
    _style(ws, headers, _BANK_WIDTHS + (12, 60))
    wrap = Alignment(wrap_text=True, vertical="top")
    for test in Test.query.filter_by(teacher_id=teacher_id).order_by(Test.created_at, Test.id):
        for q in Question.query.filter_by(test_id=test.id).order_by(Question.order_index, Question.id):
            ws.append([
                test.title, test.test_number or "", test.chapter or "", test.subject or "",
                q.q_number, q.text or "", q.option_a or "", q.option_b or "",
                q.option_c or "", q.option_d or "", q.correct_answer,
                q.marks, q.negative_marks, q.explanation or "", test.status,
                q.ai_explanation or "",
            ])
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = wrap
    return _to_bytes(wb)


def build_bank_template():
    wb = Workbook()
    ws = wb.active
    ws.title = "Question Bank"
    _style(ws, BANK_HEADERS, _BANK_WIDTHS)
    ws.append(["Business Economics - Chapter 1 Test 1", "T01", "Nature and Scope", "CA Foundation Economics", 1,
               "Human wants are _____ in response to satisfying their wants.", "Limited", "Unlimited", "Scarce", "Multiple",
               "B", 1, 0.25, "Answer: (B) Unlimited. Wants keep growing as soon as one is satisfied. Why not the others: ..."])
    ws.append(["Business Economics - Chapter 1 Test 1", "T01", "Nature and Scope", "CA Foundation Economics", 2,
               "Which of these is a factor of production?", "Demand", "Supply", "Land", "Price", "C", 1, 0.25, ""])
    return _to_bytes(wb)


def export_change_log(log):
    wb = Workbook()
    ws = wb.active
    ws.title = "Changes"
    _style(ws, ["Test", "Q No", "What changed", "Before", "After"], (34, 8, 44, 55, 55))
    wrap = Alignment(wrap_text=True, vertical="top")
    for entry in log:
        ws.append(list(entry))
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = wrap
    return _to_bytes(wb)
