"""Excel import/export helpers built on openpyxl.

Kept dependency-light and explicit: no pandas, just openpyxl reading/writing
worksheets directly. Import validates every row and returns a clear list of
errors before touching the database.
"""
import io
from datetime import datetime

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

HEADER_FILL = PatternFill(start_color="1E3A5F", end_color="1E3A5F", fill_type="solid")
HEADER_FONT = Font(color="FFFFFF", bold=True)

IMPORT_HEADERS = [
    "Question No", "Question", "Option A", "Option B", "Option C", "Option D",
    "Correct Answer", "Marks", "Negative Marks", "Explanation",
]

ROSTER_IMPORT_HEADERS = ["Roll Number", "Name", "Batch", "Registration Number", "Email", "Mobile"]


def _style_header(ws, headers):
    for col, title in enumerate(headers, start=1):
        cell = ws.cell(row=1, column=col, value=title)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for col in range(1, len(headers) + 1):
        ws.column_dimensions[get_column_letter(col)].width = 20


def build_question_import_template():
    wb = Workbook()
    ws = wb.active
    ws.title = "Questions"
    _style_header(ws, IMPORT_HEADERS)
    sample = [1, "Which of the following is a factor of production?", "Demand", "Supply", "Land", "Price", "C", 1, 0.25, ""]
    ws.append(sample)
    ws.column_dimensions["B"].width = 45
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf


def parse_question_import(file_stream, default_marks=1.0, default_negative=0.0):
    """Returns (rows, errors). rows is a list of dicts ready for DB insertion."""
    try:
        wb = load_workbook(file_stream, data_only=True)
    except Exception as exc:
        return [], [f"Could not read file: {exc}"]

    ws = wb.active
    rows = []
    errors = []

    header_row = [c.value for c in next(ws.iter_rows(min_row=1, max_row=1))]
    normalized = [str(h).strip().lower() if h else "" for h in header_row]

    def col_index(*names):
        for name in names:
            if name in normalized:
                return normalized.index(name)
        return None

    idx_qno = col_index("question no", "question number", "q. no.", "q no")
    idx_text = col_index("question", "question text")
    idx_a = col_index("option a")
    idx_b = col_index("option b")
    idx_c = col_index("option c")
    idx_d = col_index("option d")
    idx_correct = col_index("correct answer", "answer")
    idx_marks = col_index("marks")
    idx_neg = col_index("negative marks", "negative mark")
    idx_expl = col_index("explanation", "solution")

    if idx_correct is None:
        errors.append("Missing required column: 'Correct Answer'")
        return [], errors

    for i, row in enumerate(ws.iter_rows(min_row=2), start=2):
        values = [c.value for c in row]
        if all(v is None or str(v).strip() == "" for v in values):
            continue

        def get(idx):
            if idx is None or idx >= len(values):
                return None
            v = values[idx]
            return str(v).strip() if v is not None else None

        qno_raw = get(idx_qno)
        correct = (get(idx_correct) or "").upper()
        marks_raw = get(idx_marks)
        neg_raw = get(idx_neg)

        row_errors = []
        try:
            qno = int(float(qno_raw)) if qno_raw else i - 1
        except ValueError:
            row_errors.append(f"Row {i}: invalid Question No '{qno_raw}'")
            qno = i - 1

        if correct not in ("A", "B", "C", "D"):
            row_errors.append(f"Row {i}: Correct Answer must be one of A/B/C/D, got '{correct}'")

        try:
            marks = float(marks_raw) if marks_raw not in (None, "") else default_marks
        except ValueError:
            row_errors.append(f"Row {i}: invalid Marks '{marks_raw}'")
            marks = default_marks

        try:
            neg = float(neg_raw) if neg_raw not in (None, "") else default_negative
        except ValueError:
            row_errors.append(f"Row {i}: invalid Negative Marks '{neg_raw}'")
            neg = default_negative

        options = [get(idx_a), get(idx_b), get(idx_c), get(idx_d)]
        if all(o in (None, "") for o in options):
            row_errors.append(f"Row {i}: all four options are empty")

        if row_errors:
            errors.extend(row_errors)
            continue

        rows.append({
            "q_number": qno,
            "text": get(idx_text) or "",
            "option_a": options[0] or "",
            "option_b": options[1] or "",
            "option_c": options[2] or "",
            "option_d": options[3] or "",
            "correct_answer": correct,
            "marks": marks,
            "negative_marks": neg,
            "explanation": get(idx_expl) or "",
        })

    if not rows and not errors:
        errors.append("No data rows found in the file.")

    return rows, errors


def export_questions(test):
    wb = Workbook()
    ws = wb.active
    ws.title = "Questions"
    headers = IMPORT_HEADERS
    _style_header(ws, headers)
    for q in test.questions:
        ws.append([
            q.q_number, q.text or "", q.option_a or "", q.option_b or "",
            q.option_c or "", q.option_d or "", q.correct_answer,
            q.marks, q.negative_marks, q.explanation or "",
        ])
    ws.column_dimensions["B"].width = 45
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf


def export_answer_key(test):
    wb = Workbook()
    ws = wb.active
    ws.title = "Answer Key"
    ws.merge_cells("A1:D1")
    title_cell = ws["A1"]
    title_cell.value = f"{test.title} — Answer Key"
    title_cell.font = Font(bold=True, size=14)
    title_cell.alignment = Alignment(horizontal="center")

    headers = ["Q. No.", "Correct Answer", "Marks", "Negative Marks"]
    for col, h in enumerate(headers, start=1):
        cell = ws.cell(row=3, column=col, value=h)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal="center")

    for i, q in enumerate(test.questions, start=4):
        ws.cell(row=i, column=1, value=q.q_number)
        ws.cell(row=i, column=2, value=q.correct_answer)
        ws.cell(row=i, column=3, value=q.marks)
        ws.cell(row=i, column=4, value=q.negative_marks)
        for c in range(1, 5):
            ws.cell(row=i, column=c).alignment = Alignment(horizontal="center")

    for col in "ABCD":
        ws.column_dimensions[col].width = 18

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf


def export_results(test, ranked_submissions):
    """ranked_submissions: list of (rank, submission)"""
    wb = Workbook()
    ws = wb.active
    ws.title = "Results"

    headers = [
        "Rank", "Student Name", "Roll Number", "Reg. Number", "Batch",
        "Correct", "Wrong", "Unanswered", "Score", "Max Marks", "Percentage",
        "Submitted At",
    ]
    _style_header(ws, headers)

    for rank, s in ranked_submissions:
        ws.append([
            rank, s.student_name, s.roll_number, s.reg_number, s.batch,
            s.correct_count, s.wrong_count, s.unanswered_count,
            s.score, s.max_score, s.percentage,
            s.submitted_at.strftime("%Y-%m-%d %H:%M") if s.submitted_at else "",
        ])

    for col in range(1, len(headers) + 1):
        ws.column_dimensions[get_column_letter(col)].width = 16

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf


def export_detailed_responses(test, submissions):
    """One row per student per question — analysis friendly."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Detailed Responses"

    headers = [
        "Student Name", "Roll Number", "Q. No.", "Student Answer",
        "Correct Answer", "Status", "Marks Awarded",
    ]
    _style_header(ws, headers)

    for s in submissions:
        responses = {r.question_id: r for r in s.responses}
        for q in test.questions:
            r = responses.get(q.id)
            selected = r.selected_answer if r else None
            if selected is None:
                status = "Unanswered"
            elif r.is_correct:
                status = "Correct"
            else:
                status = "Wrong"
            ws.append([
                s.student_name, s.roll_number, q.q_number,
                selected or "—", q.correct_answer, status,
                r.marks_awarded if r else 0,
            ])

    for col in range(1, len(headers) + 1):
        ws.column_dimensions[get_column_letter(col)].width = 18

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf


# ---------------------------------------------------------------- roster

def build_roster_import_template():
    wb = Workbook()
    ws = wb.active
    ws.title = "Students"
    _style_header(ws, ROSTER_IMPORT_HEADERS)
    ws.append(["ECO101", "Aarav Sharma", "Batch A", "", "aarav@example.com", "9876500000"])
    for col in "ABCDEF":
        ws.column_dimensions[col].width = 20
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf


def parse_roster_import(file_stream):
    """Returns (rows, errors). rows: list of dicts with roll_number/name/etc."""
    try:
        wb = load_workbook(file_stream, data_only=True)
    except Exception as exc:
        return [], [f"Could not read file: {exc}"]

    ws = wb.active
    rows = []
    errors = []

    header_row = [c.value for c in next(ws.iter_rows(min_row=1, max_row=1))]
    normalized = [str(h).strip().lower() if h else "" for h in header_row]

    def col_index(*names):
        for name in names:
            if name in normalized:
                return normalized.index(name)
        return None

    idx_roll = col_index("roll number", "roll no", "roll no.")
    idx_name = col_index("name", "student name")
    idx_batch = col_index("batch", "batch / division", "division")
    idx_reg = col_index("registration number", "reg number", "reg no", "reg no.")
    idx_email = col_index("email")
    idx_mobile = col_index("mobile", "mobile number", "phone")

    if idx_roll is None or idx_name is None:
        errors.append("Missing required column(s): 'Roll Number' and/or 'Name'")
        return [], errors

    seen_rolls = set()
    for i, row in enumerate(ws.iter_rows(min_row=2), start=2):
        values = [c.value for c in row]
        if all(v is None or str(v).strip() == "" for v in values):
            continue

        def get(idx):
            if idx is None or idx >= len(values):
                return None
            v = values[idx]
            return str(v).strip() if v is not None else None

        roll = get(idx_roll)
        name = get(idx_name)

        if not roll:
            errors.append(f"Row {i}: Roll Number is required")
            continue
        if not name:
            errors.append(f"Row {i}: Name is required")
            continue
        if roll in seen_rolls:
            errors.append(f"Row {i}: duplicate Roll Number '{roll}' within this file")
            continue
        seen_rolls.add(roll)

        rows.append({
            "roll_number": roll, "name": name,
            "batch": get(idx_batch) or "", "reg_number": get(idx_reg) or "",
            "email": get(idx_email) or "", "mobile": get(idx_mobile) or "",
        })

    if not rows and not errors:
        errors.append("No data rows found in the file.")

    return rows, errors


def export_roster_credentials(entries):
    """entries: list of dicts {roll_number, name, batch, password} — shown/exported
    once at creation time; passwords are never stored in plaintext afterward."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Student Logins"
    headers = ["Roll Number", "Name", "Batch", "Password"]
    _style_header(ws, headers)
    for e in entries:
        ws.append([e["roll_number"], e["name"], e.get("batch") or "", e["password"]])
    for col in "ABCD":
        ws.column_dimensions[col].width = 22
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf
